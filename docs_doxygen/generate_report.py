"""Ответы на задание: Doxygen, Sphinx, docstrings."""

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH


def main():
    doc = Document()
    doc.styles["Normal"].font.name = "Times New Roman"
    doc.styles["Normal"].font.size = Pt(12)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Задание\nDoxygen, Sphinx и автоматическая документация")
    r.bold = True
    r.font.size = Pt(14)
    doc.add_paragraph("Студент: _________________________")
    doc.add_paragraph()

    doc.add_heading("1. Что делает Doxygen или Sphinx?", level=1)
    doc.add_paragraph(
        "Doxygen и Sphinx — инструменты для автоматической генерации документации из исходного кода."
    )
    doc.add_paragraph("Doxygen: создаёт HTML, PDF, CHM из комментариев в C++, Java, Python и др.")
    doc.add_paragraph("Sphinx: создаёт красивую документацию (часто HTML) из docstrings Python и файлов reStructuredText/Markdown.")
    doc.add_paragraph("Оба извлекают описания функций, классов, параметров и собирают их в единый справочник.")

    doc.add_heading("2. Из чего чаще всего генерируется документация?", level=1)
    doc.add_paragraph("• Комментарии и docstrings в коде")
    doc.add_paragraph("• Специальные разметки (Javadoc, reStructuredText, Markdown)")
    doc.add_paragraph("• Отдельные .md / .rst файлы с описанием API")
    doc.add_paragraph("• Аннотации типов и сигнатуры функций")

    doc.add_heading("3. Что такое docstring?", level=1)
    doc.add_paragraph(
        "Docstring — строка документации внутри функции, класса или модуля. "
        "В Python пишется в тройных кавычках сразу после объявления:"
    )
    doc.add_paragraph('def add(a, b):\n    """Складывает два числа."""\n    return a + b')

    doc.add_heading("4. Что произойдёт, если не писать docstrings?", level=1)
    doc.add_paragraph("• Автодокументация будет пустой или неполной")
    doc.add_paragraph("• Другим разработчикам сложнее понять назначение кода")
    doc.add_paragraph("• Придётся вручную описывать API в отдельных файлах")
    doc.add_paragraph("• Увеличивается время на разбор чужого кода и онбординг новых сотрудников")

    doc.add_heading("5. Чем удобна автоматическая документация для команды?", level=1)
    doc.add_paragraph("• Документация обновляется вместе с кодом — меньше расхождений")
    doc.add_paragraph("• Единый формат и структура справочника")
    doc.add_paragraph("• Экономия времени — не нужно писать всё вручную в Word")
    doc.add_paragraph("• Быстрый поиск по функциям и классам")
    doc.add_paragraph("• Удобно для новых участников команды и тестировщиков")

    doc.save("/workspace/docs_doxygen/Ответы_Doxygen_Sphinx.docx")
    print("OK")


if __name__ == "__main__":
    main()
